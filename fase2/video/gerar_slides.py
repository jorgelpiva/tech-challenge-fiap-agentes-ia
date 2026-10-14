"""Gera os slides de apoio do vídeo da Fase 2 (16:9) com notas do apresentador.

Uso (na raiz do repo):  python fase2/video/gerar_slides.py
Saída: fase2/video/Apresentacao_Fase2_Olist.pptx
As falas (notas) seguem o ROTEIRO_VIDEO_FASE2.md; se mudar um, mude o outro.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

AQUI = Path(__file__).resolve().parent
F2 = AQUI.parent
NAVY, AZUL, LARANJA, CINZA, INK = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x0B, 0x0B, 0x0B)
FUNDO_SUAVE = RGBColor(0xF6, 0xF6, 0xF4)
AG = {'LogiPredict': RGBColor(0x2A, 0x78, 0xD6), 'SentimentDesk': RGBColor(0x62, 0x50, 0xD6), 'WinBack': RGBColor(0x00, 0x83, 0x00)}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BRANCO = prs.slide_layouts[6]
n_slide = [0]


def texto(sl, x, y, w, h, conteudo, tam=18, cor=INK, negrito=False, alinh=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = ancora
    linhas = conteudo if isinstance(conteudo, list) else [conteudo]
    for k, ln in enumerate(linhas):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = alinh
        partes = ln.split('**')
        for j, pt in enumerate(partes):
            if not pt:
                continue
            r = p.add_run(); r.text = pt; r.font.size = Pt(tam); r.font.color.rgb = cor
            r.font.bold = negrito or (j % 2 == 1); r.font.name = 'Calibri'
        p.space_after = Pt(6)
    return tb


def novo(titulo, sub=None, notas=''):
    sl = prs.slides.add_slide(BRANCO); n_slide[0] += 1
    barra = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H); barra.fill.solid(); barra.fill.fore_color.rgb = LARANJA; barra.line.fill.background()
    texto(sl, Inches(0.55), Inches(0.35), Inches(12.2), Inches(0.8), titulo, 30, NAVY, True)
    if sub:
        texto(sl, Inches(0.55), Inches(1.05), Inches(12.2), Inches(0.5), sub, 16, CINZA)
    texto(sl, Inches(0.55), Inches(7.0), Inches(9), Inches(0.35), 'Olist Intelligent Marketplace · Tech Challenge Fase 2', 10, CINZA)
    texto(sl, Inches(12.2), Inches(7.0), Inches(0.8), Inches(0.35), str(n_slide[0]), 10, CINZA, alinh=PP_ALIGN.RIGHT)
    sl.notes_slide.notes_text_frame.text = notas
    return sl


def imagem(sl, caminho, x, y, w=None, h=None):
    return sl.shapes.add_picture(str(caminho), x, y, width=w, height=h)


def numero(sl, x, y, valor, rotulo, cor=LARANJA):
    texto(sl, x, y, Inches(3.6), Inches(0.9), valor, 40, cor, True)
    texto(sl, x, y + Inches(0.85), Inches(3.6), Inches(0.9), rotulo, 14, CINZA)


def card(sl, x, y, w, h, nome, linhas):
    cor = AG[nome]
    c = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h); c.adjustments[0] = 0.05
    c.fill.solid(); c.fill.fore_color.rgb = FUNDO_SUAVE; c.line.color.rgb = cor; c.line.width = Pt(2)
    topo = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.12)); topo.fill.solid(); topo.fill.fore_color.rgb = cor; topo.line.fill.background()
    texto(sl, x + Inches(0.2), y + Inches(0.25), w - Inches(0.4), Inches(0.6), nome, 22, cor, True)
    texto(sl, x + Inches(0.2), y + Inches(0.9), w - Inches(0.4), h - Inches(1.0), linhas, 17, INK)


# 1 · capa
sl = prs.slides.add_slide(BRANCO); n_slide[0] += 1
fundo = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H); fundo.fill.solid(); fundo.fill.fore_color.rgb = NAVY; fundo.line.fill.background()
texto(sl, Inches(0.8), Inches(1.6), Inches(11), Inches(0.5), 'TECH CHALLENGE · FASE 2', 18, RGBColor(0xF5, 0xA8, 0x86), True)
texto(sl, Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.6), 'Do diagnóstico à operação', 54, RGBColor(0xFF, 0xFF, 0xFF), True)
texto(sl, Inches(0.8), Inches(3.4), Inches(11.5), Inches(1), 'Como três agentes de IA vão proteger a jornada do pedido na Olist', 24, RGBColor(0xDB, 0xE8, 0xF9))
texto(sl, Inches(0.8), Inches(5.6), Inches(11.5), Inches(1), ['Leonardo Granjeiro · Jorge Leandro Piva · Caio Sousa · Lucas Vinicius Oliveira Mendes', 'Pós-graduação em IA Agêntica · FIAP'], 14, RGBColor(0xDB, 0xE8, 0xF9))
sl.notes_slide.notes_text_frame.text = ('[Jorge · 0:00–0:20] Na Fase 1 mostramos onde a Olist perde valor. A diretoria aprovou a continuidade, '
    'e hoje respondemos à pergunta seguinte: como os agentes vão funcionar dentro da operação e gerar valor real.')

# 2 · evolução
sl = novo('O que mudou da Fase 1 para a Fase 2', 'De seis ideias de agentes para três agentes especializados, com operação desenhada',
          notas='[Jorge · 0:20–0:50] Ouvimos a avaliação da Fase 1: faltava mostrar de onde vinham os achados, os agentes estavam amplos demais e havia objetivos demais. '
                'Respondemos com três mudanças: cada achado agora tem uma cadeia de evidência, priorizamos três agentes e cada agente ganhou um objetivo único, gatilhos e uma lista fechada de ações.')
linhas_ev = [('Fase 1', '6 agentes, objetivos amplos', 'Achados sem caminho explícito até a decisão', 'Prompts genéricos com exemplos inventados'),
             ('Fase 2', '3 agentes na jornada do pedido', 'Cadeia de evidência: dado → método → decisão', 'Prompts com ações fechadas e casos reais')]
for k, (fase, *itens) in enumerate(linhas_ev):
    x = Inches(0.6 + k * 6.3)
    box = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(5.9), Inches(4.6)); box.adjustments[0] = 0.04
    box.fill.solid(); box.fill.fore_color.rgb = FUNDO_SUAVE if k == 0 else RGBColor(0xFD, 0xF0, 0xEA)
    box.line.color.rgb = CINZA if k == 0 else LARANJA; box.line.width = Pt(1.5 if k == 0 else 2.5)
    texto(sl, x + Inches(0.3), Inches(2.0), Inches(5.3), Inches(0.7), fase, 26, CINZA if k == 0 else LARANJA, True)
    texto(sl, x + Inches(0.3), Inches(2.9), Inches(5.3), Inches(3.4), ['• ' + t for t in itens], 19, INK)

# 3 · evidência
sl = novo('O problema nasce na jornada do pedido', notas='[Leonardo · 0:50–1:20] Os dados mostram por onde começar. Quando o pedido atrasa, 54% dos clientes dão nota 1 ou 2, contra 9% no prazo. '
          'E um quarto dos atrasos começa antes da transportadora, com o seller postando depois do limite: aí o atraso sobe de 6,5% para 24,4%. Quem teve a primeira compra atrasada volta menos.')
imagem(sl, F2 / 'img' / 'f2_evidencia_atraso.png', Inches(0.5), Inches(1.4), w=Inches(8.6))
numero(sl, Inches(9.4), Inches(1.6), '54% vs 9%', 'notas 1–2 com atraso vs no prazo')
numero(sl, Inches(9.4), Inches(3.3), '27%', 'dos atrasos começam na postagem do seller')
numero(sl, Inches(9.4), Inches(5.0), '54%', 'das 2ªs compras acontecem em até 90 dias', AZUL)

# 4 · priorização
sl = novo('Priorização: menos agentes, mais perto do negócio', notas='[Leonardo · 1:20–1:40] Pontuamos os seis agentes em impacto, proximidade do core, evidência, viabilidade e risco. '
          'Três ficaram claramente à frente, e os três atuam em sequência na mesma jornada. SellerCare foi absorvido, SmartFreight foi para o backlog e o Maestro fica para a fase de orquestração.')
imagem(sl, F2 / 'img' / 'f2_priorizacao.png', Inches(0.5), Inches(1.3), h=Inches(5.6))
texto(sl, Inches(9.5), Inches(1.6), Inches(3.6), Inches(5), ['**Detalhados na Fase 2**', 'LogiPredict · 4,75', 'SentimentDesk · 4,60', 'WinBack · 4,20', '',
      '**Fora do escopo agora**', 'SellerCare → absorvido', 'SmartFreight → backlog', 'Maestro → Fase 4'], 18, INK)

# 5 · os três agentes
sl = novo('Três agentes, um papel cada', 'Cada um com gatilho, lista fechada de ações e humano nas decisões de impacto',
          notas='[Caio · 1:40–2:10] O LogiPredict é o guardião do prazo: age no limite de postagem, na postagem e no D-3. O SentimentDesk faz a triagem de toda avaliação 1 a 3, '
                'responde e leva o problema a quem resolve. O WinBack trabalha só a janela da segunda compra, de 7 a 90 dias. Nenhum deles concede reembolso, troca transportadora ou dá incentivo alto sozinho.')
cw = Inches(4.0)
dados_cards = [('LogiPredict', ['**Guardião do prazo**', 'Gatilho: limite de postagem, postagem, D-3', 'Ações: lembrete, escalar, avisar cliente', 'Humano: tratativa com transportadora', 'KPI: atraso (8,11%)']),
               ('SentimentDesk', ['**Triagem de avaliações 1–3**', 'Gatilho: avaliação publicada', 'Ações: classificar, responder, rotear', 'Humano: casos críticos e compensação', 'KPI: tempo até 1ª resposta']),
               ('WinBack', ['**Segunda compra**', 'Gatilho: D+7 após a 1ª entrega', 'Ações: pausar, mensagem, incentivo ≤ 10%', 'Humano: incentivo > 10%', 'KPI: conversão 90d (1,13%)'])]
for k, (nome, lin) in enumerate(dados_cards):
    card(sl, Inches(0.55) + k * (cw + Inches(0.25)), Inches(1.75), cw, Inches(4.9), nome, lin)

# 6-8 · fluxos
falas_fluxo = {
    'logipredict': '[Caio · 2:10–2:35] No LogiPredict, três decisões: o seller vai estourar o limite? O risco passou do limiar? A nova data passa do prazo? Só a última autoriza falar com o cliente; o analista de operações decide a tratativa.',
    'sentimentdesk': '[Lucas · 2:35–2:55] No SentimentDesk, o agente classifica com texto e dados do pedido. Confiança baixa ou caso crítico vai para o analista antes da resposta; motivo de entrega vira sinal para o LogiPredict.',
    'winback': '[Lucas · 2:55–3:10] O WinBack checa consentimento e intervalo e pausa a oferta se a primeira experiência foi ruim. Incentivo acima de 10% precisa de aprovação, e 10% dos clientes ficam como grupo controle.',
}
for nome_ag, arq in (('LogiPredict', 'logipredict'), ('SentimentDesk', 'sentimentdesk'), ('WinBack', 'winback')):
    sl = novo(f'Como opera: {nome_ag}', notas=falas_fluxo[arq])
    imagem(sl, F2 / 'fluxos' / f'fluxo_{arq}.png', Inches(0.5), Inches(1.15), w=Inches(12.3))

# 9 · exemplo de interação
sl = novo('Exemplo de interação real: SentimentDesk', 'Avaliação real do dataset processada pelo prompt v2',
          notas='[Lucas · 3:10–3:40] Um caso real. O cliente deu nota 1: "Só chegou a mochila, estou aguardando a lancheira". Uma regra de palavras-chave classificaria como "não recebido". '
                'O agente cruza com o pedido: 2 itens, entregue no prazo. Classifica como item faltando, responsabiliza o seller, abre ticket e responde ao cliente sem prometer o que não pode cumprir.')
texto(sl, Inches(0.6), Inches(1.7), Inches(5.6), Inches(0.5), 'ENTRADA', 14, CINZA, True)
texto(sl, Inches(0.6), Inches(2.1), Inches(5.6), Inches(4.5), ['Nota **1**', '"Só chegou a mochila, estou aguardando a lancheira... Quando chegar a lancheira dou mais estrelas."', '',
      'Pedido: **2 itens**, 1 seller, papelaria', 'Prazo 12/01 · entregue **03/01 (no prazo)**'], 17, INK)
texto(sl, Inches(6.8), Inches(1.7), Inches(6), Inches(0.5), 'SAÍDA DO AGENTE', 14, LARANJA, True)
texto(sl, Inches(6.8), Inches(2.1), Inches(6), Inches(4.8), ['Motivo: **PEDIDO_ITEM_FALTANDO** (0,90)', 'Severidade: ALTA · Responsável: **SELLER**', 'Roteamento: ticket ao seller', '',
      '"Sentimos muito: seu pedido deveria ter chegado completo. Já abrimos uma solicitação para localizar a lancheira. Em até 2 dias úteis você recebe nosso retorno."'], 17, INK)

# 10 · jornada
sl = novo('Os três juntos: o pedido que ia atrasar', 'Pedido real: MG → MA, postado 4 dias após o limite, entregue 10 dias depois do prazo',
          notas='[Leonardo · 3:40–4:10] Juntando tudo num pedido real: o LogiPredict lembra o seller, detecta o risco 19 dias antes do prazo e avisa o cliente com nova data. '
                'Na avaliação nota 1, o SentimentDesk responde e o WinBack pausa a oferta. Resolvido o caso, a oferta de segunda compra volta, no momento certo.')
imagem(sl, F2 / 'fluxos' / 'jornada_multiagente.png', Inches(0.5), Inches(1.55), w=Inches(12.3))

# 11 · valor
sl = novo('Valor esperado (cenário base, 12 meses, escala do dataset)', 'Premissas conservadoras; validação com grupo controle no piloto',
          notas='[Jorge · 4:10–4:35] No cenário base: 670 atrasos evitados e cerca de 300 avaliações negativas a menos, 1.229 horas de triagem poupadas e 358 segundas compras adicionais. '
                'O ganho maior vem da combinação: menos atraso gera menos detrator e preserva a recompra.')
numero(sl, Inches(0.7), Inches(1.9), '670', 'atrasos evitados · ~300 notas 1–2 a menos (LogiPredict)', AG['LogiPredict'])
numero(sl, Inches(4.9), Inches(1.9), '1.229 h', 'de triagem poupadas · R$ 55 mil (SentimentDesk)', AG['SentimentDesk'])
numero(sl, Inches(9.1), Inches(1.9), '358', '2ªs compras adicionais · R$ 53 mil (WinBack)', AG['WinBack'])
texto(sl, Inches(0.7), Inches(4.4), Inches(12), Inches(2.2), ['**Metas do piloto (90 dias):** −1 p.p. de atraso nas rotas do Nordeste · 90% das avaliações respondidas em < 24h · ≥ 85% de classificação correta · falso positivo do aviso ≤ 20%',
      '**Governança:** agente executa o que é reversível; pessoas decidem compensação, transportadora e incentivo acima do teto.'], 17, INK)

# 12 · pedido
sl = novo('Nossa recomendação', notas='[Jorge · 4:35–5:00] Pedimos a aprovação de um piloto de 90 dias com LogiPredict e SentimentDesk, focado nas rotas do Nordeste, com grupo controle e as metas que mostramos. '
          'Não propomos mais agentes; propomos agentes melhores, começando pela origem do problema. Obrigado.')
texto(sl, Inches(0.7), Inches(1.6), Inches(12), Inches(1.2), 'Aprovar o piloto da **Wave 1 (LogiPredict + SentimentDesk)** por 90 dias', 32, NAVY)
texto(sl, Inches(0.7), Inches(3.0), Inches(12), Inches(3), ['• Dados que a Olist já possui; nenhuma ação irreversível automática',
      '• Foco nas rotas do Nordeste (15–24% de atraso), com grupo controle',
      '• Wave 2: WinBack · Wave 3: SellerCare e SmartFreight · Fase 4: orquestração (Maestro)',
      '• Próximo passo (Fase 3): protótipo low-code com os casos reais deste estudo'], 20, INK)
texto(sl, Inches(0.7), Inches(5.9), Inches(12), Inches(0.8), '"Não propomos mais agentes; propomos agentes melhores."', 22, LARANJA, True)

saida = AQUI / 'Apresentacao_Fase2_Olist.pptx'
prs.save(saida)
print('ok', saida, n_slide[0], 'slides')
